import streamlit as st
import datetime
import os
from telegram import Bot
from fpdf import FPDF
from docx import Document
import asyncio
from parser import parse_all

st.set_page_config(page_title="News Monitoring Agent", layout="wide")
st.title("📰 News Monitoring Agent")

# Telegram init
TOKEN = "7908738336:AAEpY4KQHtV94Wa65qd2xsy7d_Xz1HwP4_A"
bot = Bot(token=TOKEN)

async def get_chat_id():
    updates = await bot.get_updates()
    if updates:
        chat_id = updates[-1].message.chat.id
        return chat_id
    return None

# Форма поиска
with st.form("search_form"):
    keyword = st.text_input("🔍 Ключевые слова для поиска", "Україна")
    start_date = st.date_input("📅 Начальная дата", datetime.date.today())
    end_date = st.date_input("📅 Конечная дата", datetime.date.today())
    email_to = st.text_input("📧 Email для отправки отчёта", "your@email.com")
    submitted = st.form_submit_button("Запустить поиск")

if submitted:
    st.success(f"🔎 Поиск новостей по запросу: {keyword}")
    results = parse_all(keyword)

    st.write("## Найденные новости")
    for res in results:
        st.markdown(f"**{res['title']}** — {res['date']}")
        st.markdown(res['summary'])
        st.markdown(f"[Читать далее]({res['url']})")

    # PDF
    pdf = FPDF()
    pdf.add_page()
    pdf.add_font("DejaVu", "", "DejaVuSans.ttf", uni=True)
    pdf.set_font("DejaVu", size=12)
    pdf.cell(200, 10, txt="News Report", ln=True, align='C')
    for res in results:
        text = f"{res['date']} — {res['title']}\n{res['summary']}\n{res['url']}\n"
        pdf.multi_cell(0, 10, text)
    pdf.output("report.pdf")

    # Word
    doc = Document()
    doc.add_heading("News Report", 0)
    for res in results:
        doc.add_heading(res['title'], level=1)
        doc.add_paragraph(f"{res['date']}\n{res['summary']}\n{res['url']}")
    doc.save("report.docx")

    st.success("✅ Отчёты PDF и DOCX созданы")

    # Отправка Telegram
    async def send_to_telegram():
        try:
            chat_id = await get_chat_id()
            if chat_id:
                await bot.send_message(chat_id=chat_id, text="🗞 Ваш отчёт готов! 📄")
                await bot.send_document(chat_id=chat_id, document=open("report.pdf", "rb"))
                await bot.send_document(chat_id=chat_id, document=open("report.docx", "rb"))
                st.success(f"📤 Отчёты отправлены в Telegram чат {chat_id}")
            else:
                st.warning("❗ Не удалось определить chat_id. Напиши /start боту.")
        except Exception as e:
            st.error(f"Ошибка Telegram: {e}")

    asyncio.run(send_to_telegram())
